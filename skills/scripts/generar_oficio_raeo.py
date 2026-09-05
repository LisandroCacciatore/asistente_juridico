# -*- coding: utf-8 -*-
# ============================================================
#  GENERAR_OFICIO_RAEO.PY  — Estudio Jurídico Segovia
#  UN PDF: págs. 1-2 OFICIO al RAEO (formulario) + pág. 3 ESCRITO CARGO.
#  Con --word: además un .docx SOLO del oficio (editable para el mail al juzgado).
#  Formato: cuerpo justificado, títulos centrados.
#  Uso: python3 generar_oficio_raeo.py datos.json CARPETA_SALIDA [--word]
#  Requiere: reportlab (PDF); python-docx (solo con --word).
# ============================================================
import os, re, sys, json

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

ESTILO_TITULO = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=11,
    alignment=TA_CENTER, spaceAfter=4, leading=14)
ESTILO_SECCION = ParagraphStyle("seccion", fontName="Helvetica-Bold", fontSize=10.5,
    alignment=TA_LEFT, spaceBefore=10, spaceAfter=4, leading=14)
ESTILO_CAMPO = ParagraphStyle("campo", fontName="Helvetica", fontSize=10.5,
    alignment=TA_LEFT, spaceAfter=2, leading=15)
ESTILO_CUERPO = ParagraphStyle("cuerpo", fontName="Helvetica", fontSize=10.5,
    alignment=TA_JUSTIFY, spaceAfter=8, leading=15)
ESTILO_CUERPO_ESCRITO = ParagraphStyle("cuerpo_escrito", fontName="Helvetica", fontSize=11,
    alignment=TA_JUSTIFY, spaceAfter=12, leading=17)


def _esc(t):
    if t is None:
        return ""
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _g(d, k, default=""):
    v = d.get(k)
    return v if v not in (None, "") else default


def _doc(ruta):
    return SimpleDocTemplate(ruta, pagesize=A4,
        leftMargin=2.5 * cm, rightMargin=2.5 * cm, topMargin=2.5 * cm, bottomMargin=2.5 * cm,
        title="Oficio al RAEO", author="Estudio Jurídico Segovia")


def _caratula_completa(d):
    car = _g(d, "caratula")
    expte = _g(d, "expediente_nro")
    if expte and expte not in car:
        return f"{car} {expte}".strip()
    return car


def _juzgado_desc(d):
    jd = _g(d, "juzgado_desc")
    if jd:
        return jd
    nom = _g(d, "nominacion", "____")
    ciudad = _g(d, "juzgado_localidad") or _g(d, "localidad", "Rosario")
    return f"Juzgado de Primera Instancia de Distrito en lo Laboral de la {nom} Nominación de {ciudad}"


def _campo(label, valor):
    return Paragraph(f"{_esc(label)}: {_esc(valor)}", ESTILO_CAMPO)


def elementos_oficio(d):
    e = []
    e.append(Paragraph("REGISTRO DE PROCESOS UNIVERSALES", ESTILO_TITULO))
    e.append(Paragraph("Y DE ACCIDENTES Y ENFERMEDADES", ESTILO_TITULO))
    e.append(Paragraph("OCUPACIONALES", ESTILO_TITULO))
    e.append(Spacer(1, 14))

    e.append(Paragraph("Señor", ESTILO_CAMPO))
    e.append(Paragraph("Nro.: ……………………………………", ESTILO_CAMPO))
    e.append(Paragraph("Funcionario a cargo del", ESTILO_CAMPO))
    e.append(Paragraph("Registro de Procesos Universales y de", ESTILO_CAMPO))
    e.append(Paragraph("Accidentes y Enfermedades Ocupacionales", ESTILO_CAMPO))
    e.append(Paragraph("Fecha: ………………………………", ESTILO_CAMPO))
    e.append(Paragraph("S / D", ESTILO_CAMPO))
    e.append(Spacer(1, 10))

    e.append(Paragraph(
        "Quien suscribe informa a Ud. que se han iniciado las siguientes actuaciones "
        "sobre accidentes y enfermedades ocupacionales:", ESTILO_CUERPO))

    e.append(Paragraph("I - Datos del trabajador accidentado o enfermo", ESTILO_SECCION))
    e.append(_campo("Apellido", _g(d, "apellido")))
    e.append(_campo("Nombres", _g(d, "nombres")))
    e.append(_campo("Actividad", _g(d, "actividad")))
    e.append(_campo("Antigüedad en ese empleo", _g(d, "antiguedad")))
    e.append(_campo("Edad", _g(d, "edad")))
    e.append(Paragraph(
        f"Doc. Tipo: {_esc(_g(d, 'doc_tipo', 'DNI'))}&nbsp;&nbsp;Nro.: {_esc(_g(d, 'doc_nro'))}"
        f"&nbsp;&nbsp;Estado Civil: {_esc(_g(d, 'estado_civil'))}", ESTILO_CAMPO))
    e.append(Paragraph(
        f"Domicilio: {_esc(_g(d, 'domicilio'))}&nbsp;&nbsp;Localidad: {_esc(_g(d, 'localidad'))}",
        ESTILO_CAMPO))

    e.append(Paragraph("II - Datos del accidente y/o enfermedad", ESTILO_SECCION))
    e.append(Paragraph(
        f"Fecha: {_esc(_g(d, 'fecha_accidente'))}&nbsp;&nbsp;"
        f"Localidad donde ocurrió: {_esc(_g(d, 'localidad_ocurrio'))}", ESTILO_CAMPO))
    e.append(Paragraph(
        f"Circunstancias de lugar y modo del accidente y/o enfermedad: "
        f"{_esc(_g(d, 'circunstancias'))}", ESTILO_CUERPO))
    e.append(Paragraph(
        f"Lesión y/o enfermedad denunciada: {_esc(_g(d, 'lesion'))}", ESTILO_CAMPO))

    e.append(Paragraph("III - Datos del reclamo", ESTILO_SECCION))
    e.append(_campo("Concepto reclamado (objeto de la demanda)", _g(d, "concepto")))
    e.append(_campo("Porcentaje de Incapacidad reclamada", _g(d, "porcentaje_incapacidad")))
    e.append(_campo("Monto reclamado", _g(d, "monto_reclamado")))

    e.append(Paragraph("IV - Datos de los responsables demandados", ESTILO_SECCION))
    e.append(_campo("Nombres del empleador y/o responsables", _g(d, "empleador_nombre")))
    e.append(Paragraph(
        f"Domicilio: {_esc(_g(d, 'empleador_domicilio'))}&nbsp;&nbsp;"
        f"Localidad: {_esc(_g(d, 'empleador_localidad'))}", ESTILO_CAMPO))
    e.append(_campo("Ramo o actividad", _g(d, "ramo")))
    if _g(d, "otros_responsables"):
        e.append(_campo("Otros responsables o demandados", _g(d, "otros_responsables")))

    e.append(Paragraph("V - Datos del Proceso", ESTILO_SECCION))
    e.append(_campo("Expediente Nro.", _g(d, "expediente_nro")))
    e.append(_campo("Carátula", _g(d, "caratula")))
    e.append(_campo("Vínculo del Actor con el accidentado", _g(d, "vinculo_actor", "Es la misma persona")))
    e.append(_campo("Profesionales del actor", _g(d, "profesionales", "SANTIAGO AGUSTÍN SEGOVIA")))
    e.append(_campo("Acción interpuesta", _g(d, "accion", "Especial (Ley 27.348)")))
    e.append(Paragraph(
        f"Juzgado: {_esc(_g(d, 'juzgado'))}&nbsp;&nbsp;"
        f"Localidad: {_esc(_g(d, 'juzgado_localidad') or _g(d, 'localidad', 'Rosario'))}", ESTILO_CAMPO))
    e.append(_campo("Fecha de iniciación de la causa", _g(d, "fecha_iniciacion")))
    if _g(d, "observaciones"):
        e.append(_campo("Observaciones", _g(d, "observaciones")))

    e.append(Spacer(1, 40))
    e.append(Paragraph("FIRMA Y SELLO DEL PROFESIONAL", ESTILO_CAMPO))
    return e


def elementos_escrito(d):
    abogado = _g(d, "abogado", "Santiago A. Segovia")
    return [
        Paragraph("ACOMPAÑA OFICIO AL RAEO PARA SU DILIGENCIAMIENTO", ESTILO_TITULO),
        Spacer(1, 16),
        Paragraph("<b>Sr. Juez:</b>", ESTILO_CUERPO_ESCRITO),
        Paragraph(
            f"{_esc(abogado)}, abogado por la parte actora, en la representación acreditada "
            f"dentro de los autos caratulados: &ldquo;<b>{_esc(_caratula_completa(d))}</b>&rdquo;, "
            f"en trámite ante el {_esc(_juzgado_desc(d))}, ante V.S. me presento y digo:",
            ESTILO_CUERPO_ESCRITO),
        Paragraph(
            "Que vengo por la presente a acompañar oficio para ser diligenciado al RAEO, "
            "asimismo, se acompañó en formato editable vía mail a la casilla de correo del "
            "juzgado.", ESTILO_CUERPO_ESCRITO),
        Paragraph("Por lo expuesto a V.S. solicito:", ESTILO_CUERPO_ESCRITO),
        Paragraph("1) Tenga por acompañado proyecto de oficio al RAEO.", ESTILO_CUERPO_ESCRITO),
        Paragraph("2) Se disponga su rúbrica y libramiento.", ESTILO_CUERPO_ESCRITO),
        Spacer(1, 10),
        Paragraph("Proveer de conformidad,", ESTILO_CUERPO_ESCRITO),
        Paragraph("<b>Y SERÁ JUSTICIA.</b>", ESTILO_CUERPO_ESCRITO),
    ]


def generar_pdf(d, ruta):
    elementos = elementos_oficio(d) + [PageBreak()] + elementos_escrito(d)
    _doc(ruta).build(elementos)
    return ruta


def generar_docx_oficio(d, ruta):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Times New Roman"
    base.font.size = Pt(11)

    def par(texto="", bold=False, center=False, justify=False, space_before=0):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.space_before = Pt(space_before)
        if texto:
            r = p.add_run(texto)
            r.bold = bold
        return p

    def campo(label, valor):
        par(f"{label}: {valor}")

    par("REGISTRO DE PROCESOS UNIVERSALES", bold=True, center=True)
    par("Y DE ACCIDENTES Y ENFERMEDADES", bold=True, center=True)
    par("OCUPACIONALES", bold=True, center=True)
    par()
    par("Señor")
    par("Nro.: ……………………………………")
    par("Funcionario a cargo del")
    par("Registro de Procesos Universales y de")
    par("Accidentes y Enfermedades Ocupacionales")
    par("Fecha: ………………………………")
    par("S / D")
    par()
    par("Quien suscribe informa a Ud. que se han iniciado las siguientes actuaciones sobre "
        "accidentes y enfermedades ocupacionales:", justify=True)

    par("I - Datos del trabajador accidentado o enfermo", bold=True, space_before=8)
    campo("Apellido", _g(d, "apellido"))
    campo("Nombres", _g(d, "nombres"))
    campo("Actividad", _g(d, "actividad"))
    campo("Antigüedad en ese empleo", _g(d, "antiguedad"))
    campo("Edad", _g(d, "edad"))
    campo("Doc. Tipo", f"{_g(d, 'doc_tipo', 'DNI')}   Nro.: {_g(d, 'doc_nro')}   "
          f"Estado Civil: {_g(d, 'estado_civil')}")
    campo("Domicilio", f"{_g(d, 'domicilio')}   Localidad: {_g(d, 'localidad')}")

    par("II - Datos del accidente y/o enfermedad", bold=True, space_before=8)
    campo("Fecha", f"{_g(d, 'fecha_accidente')}   Localidad donde ocurrió: {_g(d, 'localidad_ocurrio')}")
    par(f"Circunstancias de lugar y modo del accidente y/o enfermedad: {_g(d, 'circunstancias')}",
        justify=True)
    campo("Lesión y/o enfermedad denunciada", _g(d, "lesion"))

    par("III - Datos del reclamo", bold=True, space_before=8)
    campo("Concepto reclamado (objeto de la demanda)", _g(d, "concepto"))
    campo("Porcentaje de Incapacidad reclamada", _g(d, "porcentaje_incapacidad"))
    campo("Monto reclamado", _g(d, "monto_reclamado"))

    par("IV - Datos de los responsables demandados", bold=True, space_before=8)
    campo("Nombres del empleador y/o responsables", _g(d, "empleador_nombre"))
    campo("Domicilio", f"{_g(d, 'empleador_domicilio')}   Localidad: {_g(d, 'empleador_localidad')}")
    campo("Ramo o actividad", _g(d, "ramo"))
    if _g(d, "otros_responsables"):
        campo("Otros responsables o demandados", _g(d, "otros_responsables"))

    par("V - Datos del Proceso", bold=True, space_before=8)
    campo("Expediente Nro.", _g(d, "expediente_nro"))
    campo("Carátula", _g(d, "caratula"))
    campo("Vínculo del Actor con el accidentado", _g(d, "vinculo_actor", "Es la misma persona"))
    campo("Profesionales del actor", _g(d, "profesionales", "SANTIAGO AGUSTÍN SEGOVIA"))
    campo("Acción interpuesta", _g(d, "accion", "Especial (Ley 27.348)"))
    campo("Juzgado", f"{_g(d, 'juzgado')}   Localidad: "
          f"{_g(d, 'juzgado_localidad') or _g(d, 'localidad', 'Rosario')}")
    campo("Fecha de iniciación de la causa", _g(d, "fecha_iniciacion"))
    if _g(d, "observaciones"):
        campo("Observaciones", _g(d, "observaciones"))

    par()
    par()
    par("FIRMA Y SELLO DEL PROFESIONAL")

    doc.save(ruta)
    return ruta


def _apellido(d):
    ap = _g(d, "apellido")
    if ap:
        return ap.upper().replace(" ", "-")
    m = re.match(r"^\s*([A-Za-zÁÉÍÓÚÑáéíóúñ]+)", _g(d, "caratula"))
    return m.group(1).upper() if m else "ACTOR"


def _cuij(d):
    return re.sub(r"[^0-9-]", "", _g(d, "expediente_nro")) or "SINCUIJ"


def main():
    if len(sys.argv) < 3:
        print("Uso: python3 generar_oficio_raeo.py datos.json CARPETA_SALIDA [--word]")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        d = json.load(f)
    carpeta = sys.argv[2]
    os.makedirs(carpeta, exist_ok=True)
    hacer_word = "--word" in sys.argv[3:]

    base = f"OFICIO_RAEO_{_apellido(d)}_{_cuij(d)}"
    ruta_pdf = os.path.join(carpeta, base + ".pdf")
    generar_pdf(d, ruta_pdf)
    print("PDF (oficio + escrito cargo):\n ", ruta_pdf)

    if hacer_word:
        ruta_docx = os.path.join(carpeta, base + ".docx")
        generar_docx_oficio(d, ruta_docx)
        print("Word (solo oficio):\n ", ruta_docx)


if __name__ == "__main__":
    main()
